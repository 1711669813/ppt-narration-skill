# -*- coding: utf-8 -*-
"""Build the cue-synced narrated slide video.

For every spoken cue the matching slide region is dimmed-out/highlighted, small blocks are
zoomed in, and the page turns when the narration of that page ends (+1s hold).
"""
import sys
sys.dont_write_bytecode = True
from output_paths import build_root, delivery_root, output_path, MODEL_CACHE

import asyncio, hashlib, json, os, re, subprocess, sys, math
import edge_tts, mutagen.mp3
from PIL import Image, ImageDraw

ROOT = build_root()
sys.path.insert(0, ROOT)
import importlib
CUES = importlib.import_module(os.environ.get("CUES_MODULE", "cues")).CUES
SLIDES4K = os.path.join(ROOT, "slides_png_4k")
AUD = os.path.join(ROOT, "out2", "audio")
FRM = os.path.join(ROOT, "out2", "frames")
OUT_MP4 = os.path.join(ROOT, "out2", "讲解录屏_1080p.mp4")
PPTX = os.environ.get("PPTX")
if not PPTX or not os.path.isfile(PPTX):
    sys.exit("请用环境变量 PPTX 指定存在的源 .pptx")

VOICE = os.environ.get("VOICE", "zh-CN-YunyangNeural")
RATE = os.environ.get("TTS_RATE", "+6%")
PAUSE = 1.0        # hold after the page's speech before turning the page
FPS = 30
TRANS = 12         # transition frames (x 1/30 s = 0.4 s)
DIM_BOX = 72       # darkening outside a box/zoom mark (0-255 alpha of black over the page)
DIM_LINE = 46      # lighter darkening for underline marks
ACCENT = (255, 190, 0)      # highlight border
W, H = 1920, 1080
SW, SH = 3840, 2160
MIN_VIEW_W = 0.40  # never zoom in further than 40% of the slide width (keeps text sharp)

os.makedirs(AUD, exist_ok=True)
os.makedirs(FRM, exist_ok=True)


def run(args):
    p = subprocess.run(args, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if p.returncode != 0:
        raise RuntimeError("cmd failed: %s\n%s" % (" ".join(args[:8]), (p.stderr or "")[-2000:]))
    return p.stderr or ""


# ---------------------------------------------------------------- table row rects
def table_row_rects():
    from pptx import Presentation
    prs = Presentation(PPTX)
    out = {}
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if getattr(sh, "has_table", False) and sh.has_table:
                tops, tot = [], sum(r.height for r in sh.table.rows)
                y = sh.top
                for r in sh.table.rows:
                    tops.append((y / prs.slide_height, (y + r.height) / prs.slide_height))
                    y += r.height
                out[(i, sh.name)] = (sh.left / prs.slide_width, (sh.left + sh.width) / prs.slide_width, tops)
    return out


TROW = table_row_rects()


def cue_rects(page, cue):
    rects = [list(r) for r in cue.get("r", [])]
    tr = cue.get("tr")
    if tr:
        name, idx = tr
        left, right, tops = TROW[(page, name)]
        for k in idx:
            y0, y1 = tops[k]
            rects.append([left, y0, right, y1])
    return rects


# ---------------------------------------------------------------- TTS with word boundaries
async def tts_page(page, text):
    tag = hashlib.md5((VOICE + RATE + text).encode("utf-8")).hexdigest()[:8]
    mp3 = os.path.join(AUD, "page%02d_%s.mp3" % (page, tag))
    wb = os.path.join(AUD, "page%02d_%s.words.json" % (page, tag))
    if os.path.exists(mp3) and os.path.exists(wb) and os.environ.get("FORCE_TTS") != "1":
        return mp3, json.load(open(wb, encoding="utf-8"))
    for attempt in range(4):
        try:
            comm = edge_tts.Communicate(text, VOICE, rate=RATE, boundary="WordBoundary")
            data, words = bytearray(), []
            async for chunk in comm.stream():
                if chunk["type"] == "audio":
                    data += chunk["data"]
                elif chunk["type"] == "WordBoundary":
                    words.append({"t": chunk["text"], "s": chunk["offset"] / 1e7,
                                  "d": chunk["duration"] / 1e7})
            if len(data) > 2000:
                open(mp3, "wb").write(bytes(data))
                json.dump(words, open(wb, "w", encoding="utf-8"), ensure_ascii=False)
                return mp3, words
        except Exception as e:
            print("   tts retry", page, repr(e)[:120])
            await asyncio.sleep(2)
    raise SystemExit("TTS failed page %d" % page)


def speech_end(mp3):
    err = run(["ffmpeg", "-hide_banner", "-nostats", "-i", mp3,
               "-af", "silencedetect=noise=-40dB:d=0.30", "-f", "null", "-"])
    starts = [float(m) for m in re.findall(r"silence_start:\s*([\d.]+)", err)]
    total = mutagen.mp3.MP3(mp3).info.length
    return min(max(starts), total) if starts else total


# ---------------------------------------------------------------- cue timing
def norm(s):
    return re.sub(r"[^\w\u4e00-\u9fff]", "", s)


def align(page, cues, words, sp_end):
    """cue start/end (page-local seconds) from edge-tts word boundaries; proportional fallback."""
    texts = [c["t"] for c in cues]
    full = "".join(texts)
    nfull = norm(full)
    joined = norm("".join(w["t"] for w in words))
    bounds, ok = [], False
    if joined == nfull and words:
        ok = True
        pos, wi, wpos = 0, 0, 0
        starts = []
        for w in words:  # char index in normalized text where each word starts
            starts.append(wpos)
            wpos += len(norm(w["t"]))
        for ci, t in enumerate(texts):
            nt = norm(t)
            a, b = pos, pos + len(nt)
            first = max([i for i, s in enumerate(starts) if s <= a] or [0])
            last = max([i for i, s in enumerate(starts) if s < b] or [first])
            s = words[first]["s"]
            e = words[last]["s"] + words[last]["d"]
            bounds.append([s, e])
            pos = b
    else:
        print("   [warn] page %d: word boundaries do not match text exactly -> proportional timing" % page)
        tot = len(nfull) or 1
        acc = 0
        for t in texts:
            a = acc / tot * sp_end
            acc += len(norm(t))
            bounds.append([a, acc / tot * sp_end])
    # lead-in so the highlight lands just before the words
    for i, b in enumerate(bounds):
        b[0] = max(0.0, b[0] - 0.14)
    for i in range(1, len(bounds)):
        if bounds[i][0] < bounds[i - 1][1]:
            mid = (bounds[i][0] + bounds[i - 1][1]) / 2
            bounds[i - 1][1] = bounds[i][0] = mid
    bounds[0][0] = 0.0
    bounds[-1][1] = max(bounds[-1][1], sp_end)
    return bounds, ok


# ---------------------------------------------------------------- frame rendering
def cue_style(cue):
    """mark style for one cue: 'none' (no dim, no box) | 'line' (underline) | 'box' | 'zoom'."""
    return cue.get("m") or ("zoom" if cue.get("z") == "auto" else "box")


def ts(sec):
    """seconds -> srt timestamp"""
    millis = max(0, round(sec * 1000))
    h, millis = divmod(millis, 3600000)
    m, millis = divmod(millis, 60000)
    sec, millis = divmod(millis, 1000)
    return "%02d:%02d:%02d,%03d" % (h, m, sec, millis)


def wrap2(text, width=22):
    """split one cue into at most two subtitle lines, breaking on punctuation near the middle"""
    t = text.strip()
    if len(t) <= width:
        return t
    mid = len(t) // 2
    cands = [i for i, ch in enumerate(t) if ch in "，。；、：,;" and 6 <= i <= len(t) - 6]
    cut = min(cands, key=lambda i: abs(i - mid)) + 1 if cands else mid
    return t[:cut] + "\n" + t[cut:]


def marks_of(cue, alpha):
    """(rect, alpha, style) marks for one plan cue; a 'none' cue contributes nothing."""
    st = cue.get("m", "box")
    if st == "none" or alpha <= 0:
        return []
    return [(r, alpha, st) for r in cue["r"]]


def style_dim(style):
    """darkening outside the marked block, per mark style"""
    return {"box": DIM_BOX, "zoom": DIM_BOX, "line": DIM_LINE}.get(style, 0)


def view_rect(rects, style):
    """'zoom' -> zoom onto the target; anything else -> whole page"""
    if style != "zoom" or not rects:
        return [0.0, 0.0, 1.0, 1.0]
    x0 = min(r[0] for r in rects); y0 = min(r[1] for r in rects)
    x1 = max(r[2] for r in rects); y1 = max(r[3] for r in rects)
    w, h = (x1 - x0), (y1 - y0)
    if w >= 0.45 or h >= 0.45 or w <= 0 or h <= 0:
        return [0.0, 0.0, 1.0, 1.0]
    # NOTE: in normalised page fractions a 16:9 window of a 16:9 page is a SQUARE
    # (vw*3840)/(vh*2160) == 16/9  <=>  vw == vh).  Anything else distorts the crop.
    need = min(max(max(w, h) * 2.3, MIN_VIEW_W), 1.0)
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    vx0 = min(max(cx - need / 2, 0.0), max(0.0, 1.0 - need))
    vy0 = min(max(cy - need / 2, 0.0), max(0.0, 1.0 - need))
    return [vx0, vy0, vx0 + need, vy0 + need]


def to_view(rect, view):
    """slide rect -> view-local 0..1 rect"""
    vx0, vy0, vx1, vy1 = view
    vw, vh = vx1 - vx0, vy1 - vy0
    return [(rect[0] - vx0) / vw, (rect[1] - vy0) / vh,
            (rect[2] - vx0) / vw, (rect[3] - vy0) / vh]


def render_frame(page_img, view, marks, path):
    """marks: [(rect, alpha, style)] - style in {'none','line','box','zoom'}"""
    box = (int(view[0] * SW), int(view[1] * SH), int(math.ceil(view[2] * SW)), int(math.ceil(view[3] * SH)))
    frame = page_img.crop(box).resize((W, H), Image.LANCZOS).convert("RGB")
    dim = int(round(max([style_dim(s) * a for _, a, s in marks] or [0])))
    vis = []
    for rect, alpha, style in marks:
        if alpha <= 0.03 or style == "none":
            continue
        v = to_view(rect, view)
        pad = 7
        x0 = max(0, v[0] * W - pad); y0 = max(0, v[1] * H - pad)
        x1 = min(W, v[2] * W + pad); y1 = min(H, v[3] * H + pad)
        if x1 - x0 < 8 or y1 - y0 < 8:
            continue
        vis.append((x0, y0, x1, y1, alpha, style))
    if vis and dim > 0:                      # darken everything outside the marked blocks
        mask = Image.new("L", (W, H), dim)
        md = ImageDraw.Draw(mask)
        for x0, y0, x1, y1, alpha, style in vis:
            md.rounded_rectangle([x0, y0, x1, y1], radius=max(6, min(18, int((y1 - y0) * 0.12))),
                                 fill=int(dim * (1 - alpha)))
        frame = Image.composite(Image.new("RGB", (W, H), (0, 0, 0)), frame, mask)
    ov = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(ov)
    for x0, y0, x1, y1, alpha, style in vis:
        if style == "line":                  # underline just below the marked block
            th = max(5, min(10, int((y1 - y0) * 0.05)))
            y = min(H - th, y1 + 7)
            d.rounded_rectangle([x0 - 3, y - 6, x1 + 3, y + th + 6], radius=th,
                                fill=ACCENT + (int(38 * alpha),))
            d.rounded_rectangle([x0, y, x1, y + th], radius=th // 2, fill=ACCENT + (int(255 * alpha),))
        else:                                # rounded highlight box
            r = max(6, min(18, int((y1 - y0) * 0.12)))
            d.rounded_rectangle([x0 - 4, y0 - 4, x1 + 4, y1 + 4], radius=r + 4,
                                outline=ACCENT + (int(60 * alpha),), width=9)
            d.rounded_rectangle([x0, y0, x1, y1], radius=r, outline=ACCENT + (int(235 * alpha),), width=5)
    frame = Image.alpha_composite(frame.convert("RGBA"), ov).convert("RGB")
    frame.save(path, "JPEG", quality=92, subsampling=0)


def lerp(a, b, t):
    return [a[i] + (b[i] - a[i]) * t for i in range(len(a))]


# ---------------------------------------------------------------- main build
async def main():
    pages = sorted(CUES)
    plan, frames = [], []
    nframes = 0
    for page in pages:
        cues = CUES[page]
        text = "".join(c["t"] for c in cues)
        mp3, words = await tts_page(page, text)
        sp_end = speech_end(mp3)
        bounds, ok = align(page, cues, words, sp_end)
        rects = [cue_rects(page, c) for c in cues]
        views = [view_rect(rects[i], cue_style(cues[i])) for i in range(len(cues))]
        for i, v in enumerate(views):           # geometry guards (square == 16:9, inside the page)
            if v != [0.0, 0.0, 1.0, 1.0]:
                assert abs((v[2] - v[0]) - (v[3] - v[1])) < 1e-9, \
                    "view not square (=not 16:9) page %d cue %d" % (page, i + 1)
                assert -1e-9 <= v[0] and v[2] <= 1 + 1e-9 and -1e-9 <= v[1] and v[3] <= 1 + 1e-9, \
                    "view outside page %d cue %d" % (page, i + 1)
        print("   marks:", " ".join(
            cue_style(cues[i]) + ("" if views[i] == [0.0, 0.0, 1.0, 1.0]
                                  else "@%.2fx" % (1.0 / (views[i][2] - views[i][0])))
            for i in range(len(cues))))
        page_dur = sp_end + PAUSE
        plan.append({"page": page, "mp3": mp3, "dur": page_dur, "sp_end": sp_end, "wb": ok,
                     "cues": [{"i": i, "t": cues[i]["t"], "lb": cues[i]["lb"], "r": rects[i],
                               "s": bounds[i][0], "e": bounds[i][1], "view": views[i],
                               "m": cue_style(cues[i]),
                               "zoom": views[i] != [0.0, 0.0, 1.0, 1.0]} for i in range(len(cues))]})
        print("page %2d  cues %2d  speech %6.2fs  page %6.2fs  wordboundary=%s"
              % (page, len(cues), sp_end, page_dur, "exact" if ok else "fallback"))

    total = sum(p["dur"] for p in plan)
    print("TOTAL %.1fs = %d:%02d" % (total, total // 60, total % 60))

    for p in plan:                     # absolute times
        p["start"] = sum(q["dur"] for q in plan if q["page"] < p["page"])

    # ---- frames
    step = 1.0 / FPS
    for p in plan:
        page_img = Image.open(os.path.join(SLIDES4K, "slide%02d.png" % p["page"]))
        cues = p["cues"]
        pf = []                                    # (path, duration) for this page
        for i, c in enumerate(cues):
            last = i == len(cues) - 1
            if last:
                hold = p["dur"] - c["s"]
            else:
                hold = (cues[i + 1]["s"] - c["s"]) - TRANS * step
            hold = max(round(hold / step) * step, step)
            f = os.path.join(FRM, "p%02d_c%02d_hold.jpg" % (p["page"], i))
            render_frame(page_img, c["view"], marks_of(c, 1.0), f)
            pf.append((f, hold))
            if not last:
                nxt = cues[i + 1]
                for k in range(1, TRANS + 1):
                    t = k / TRANS
                    ft = os.path.join(FRM, "p%02d_c%02d_t%02d.jpg" % (p["page"], i, k))
                    hl = marks_of(c, 1 - t) + marks_of(nxt, t)
                    render_frame(page_img, lerp(c["view"], nxt["view"], t), hl, ft)
                    pf.append((ft, step))
        page_target = round(p["dur"] / step) * step
        residual = page_target - sum(d for _, d in pf)
        path, d = pf[-1]
        pf[-1] = (path, round((d + residual) / step) * step)
        for path, d in pf:
            frames.append((path, d))
        nframes += len(pf)
        p["dur"] = sum(d for _, d in pf)
        print("  page %2d frames done (%d total, page %.2fs)" % (p["page"], nframes, p["dur"]))

    elapsed = 0.0
    for p in plan:
        p["start"] = elapsed
        elapsed += p["dur"]
    vsum = sum(d for _, d in frames)
    print("frame-count %d, sum of frame durations %.3fs, target %.3fs" % (len(frames), vsum, total))

    # ---- narration wav (page audio trimmed/padded to page length, concatenated)
    args = ["ffmpeg", "-y", "-hide_banner", "-loglevel", "error"]
    for p in plan:
        args += ["-i", p["mp3"]]
    fc = ["[%d:a]atrim=end=%.3f,apad=whole_dur=%.3f,aresample=48000,aformat=channel_layouts=stereo,"
          "asetpts=N/SR/TB[a%d]" % (i, p["dur"], p["dur"], i) for i, p in enumerate(plan)]
    fc.append("".join("[a%d]" % i for i in range(len(plan))) + "concat=n=%d:v=0:a=1[out]" % len(plan))
    wav = os.path.join(os.path.dirname(OUT_MP4), "narration.wav")
    run(args + ["-filter_complex", ";".join(fc), "-map", "[out]", "-c:a", "pcm_s16le", wav])
    print("narration wav ok")

    # ---- encode
    lst = os.path.join(os.path.dirname(OUT_MP4), "frames.txt")
    with open(lst, "w", encoding="utf-8") as f:
        for path, d in frames:
            f.write("file '%s'\nduration %.5f\n" % (path.replace("\\", "/"), d))
        f.write("file '%s'\n" % frames[-1][0].replace("\\", "/"))
    run(["ffmpeg", "-y", "-hide_banner", "-loglevel", "error", "-f", "concat", "-safe", "0", "-i", lst,
         "-i", wav, "-map", "0:v", "-map", "1:a",
         "-vf", "scale=1920:1080,setsar=1,fps=%d,format=yuv420p" % FPS,
         "-c:v", "libx264", "-preset", "medium", "-crf", "18", "-g", "60",
         "-c:a", "aac", "-b:a", "192k", "-ar", "48000", "-t", "%.3f" % vsum,
         "-movflags", "+faststart", OUT_MP4])
    print("MP4:", OUT_MP4, os.path.getsize(OUT_MP4))

    json.dump({"voice": VOICE, "rate": RATE, "fps": FPS, "pause": PAUSE, "total": vsum,
               "pages": plan}, open(os.path.join(ROOT, "timing2.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)

    if os.environ.get("SUBTITLES", "1") == "0":
        return

    # ---- sidecar subtitles (.srt), one cue per subtitle, no burn-in
    srt = os.path.join(os.path.dirname(OUT_MP4), "subtitles.srt")
    with open(srt, "w", encoding="utf-8") as f:
        n = 0
        for p in plan:
            for c in p["cues"]:
                n += 1
                f.write("%d\n%s --> %s\n%s\n\n" % (n, ts(p["start"] + c["s"]), ts(p["start"] + c["e"]),
                                                   wrap2(c["t"])))
    print("SRT:", srt, n, "lines")


if __name__ == "__main__":
    asyncio.run(main())
