"""Packaged integration test: local mock AI + real PowerPoint, TTS and FFmpeg.

Run: python tests/smoke_desktop.py path/to/packaged.exe
The mock verifies request format; it does not validate a real AI provider.
"""
import json
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import subprocess
import sys
import threading

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1] / "results" / "desktop-integration"
ROOT.mkdir(parents=True, exist_ok=True)


class Handler(BaseHTTPRequestHandler):
    requests_seen = 0

    def log_message(self, *args):
        pass

    def do_POST(self):
        data = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        assert self.path == "/v1/chat/completions"
        assert self.headers["Authorization"] == "Bearer test-placeholder-key"
        assert data["model"] == "test-vision-model"
        content = data["messages"][1]["content"]
        assert content[1]["image_url"]["url"].startswith("data:image/jpeg;base64,")
        prompt = json.loads(content[0]["text"])
        page = prompt["page"]
        type(self).requests_seen += 1
        cues = [{"text": "我们来看本页的内容。", "mark": "none", "targets": []},
                {"text": "稀疏传感可以帮助分析桥梁状态。", "mark": "zoom", "targets": [{"shape": 1}]}]
        answer = {"title": "测试页面%d" % page, "cues": cues}
        response = json.dumps({"choices": [{"message": {"content": json.dumps(answer, ensure_ascii=False)}}]}).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(response)))
        self.end_headers()
        self.wfile.write(response)


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    for page in range(1, 4):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(.7), Inches(.5), Inches(10), Inches(.8))
        title.text_frame.paragraphs[0].text = "桥梁状态分析 · 第%d页" % page
        title.text_frame.paragraphs[0].font.size = Pt(32)
        box = slide.shapes.add_textbox(Inches(4.5), Inches(3), Inches(3), Inches(1))
        box.text_frame.paragraphs[0].text = "稀疏传感与损伤识别"
        box.text_frame.paragraphs[0].font.size = Pt(22)
    ppt = ROOT / "integration-input.pptx"
    prs.save(ppt)
    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    config = {"pptx": str(ppt), "base_url": "http://127.0.0.1:%d/v1" % server.server_port,
              "api_key": "test-placeholder-key", "model": "test-vision-model", "style": "严谨学术汇报",
              "voice": "zh-CN-XiaoxiaoNeural", "rate": "+6%", "minutes": 2, "output": str(ROOT),
              "vision": True, "subtitles": True}
    try:
        command = [sys.executable, sys.argv[1], "--worker"] if sys.argv[1].endswith(".py") else [sys.argv[1], "--worker"]
        result = subprocess.run(command, input=json.dumps(config) + "\n",
                                capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=420)
    finally:
        server.shutdown()
    (ROOT / "integration.log").write_text(result.stdout + result.stderr, encoding="utf-8")
    print(result.stdout[-4500:])
    print(result.stderr[-1000:])
    assert result.returncode == 0, "Packaged pipeline failed; inspect integration.log"
    assert Handler.requests_seen == 3
    data = json.loads(next(line for line in result.stdout.splitlines() if line.startswith("@@RESULT@@"))[10:])
    assert data["ok"]
    root = Path(data["root"])
    assert (root / "讲解录屏_1080p.mp4").is_file()
    assert (root / "subtitles.srt").is_file()
    assert not any("[失败]" in line for line in result.stdout.splitlines())
    plan = json.loads((root / "timing2.json").read_text(encoding="utf-8"))
    assert len(plan["pages"]) == 3 and plan["pages"][1]["cues"][1]["m"] == "zoom"
    assert plan["pages"][1]["wb"], "TTS word timestamps were not aligned"
    print("PACKAGED INTEGRATION PASSED:", root)


if __name__ == "__main__":
    main()
